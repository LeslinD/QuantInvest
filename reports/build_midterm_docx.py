from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "reports" / "体育大事件事件驱动量化投资系统_中期报告.docx"
RESULTS = ROOT / "outputs" / "results"


BLUE = "2E74B5"
DARK = "1F4D78"
MUTED = "666666"
LIGHT = "F2F4F7"
BORDER = "DADCE0"
CALLOUT = "F4F6F9"


def load_data():
    summary = json.loads((RESULTS / "run_summary_2026-05-26.json").read_text(encoding="utf-8"))
    orders = pd.read_csv(RESULTS / "paper_orders_2026-05-27.csv")
    trades = pd.read_csv(RESULTS / "backtest_trades_2026-05-26.csv")
    wf = pd.read_csv(RESULTS / "walk_forward_trades_2026-05-26.csv")
    car = pd.read_csv(RESULTS / "event_study_summary_2026-05-26.csv")
    hp = pd.read_csv(RESULTS / "hyperparam_search_2026-05-26.csv")
    scores = pd.read_csv(RESULTS / "paper_scores_2026-05-26.csv")
    diagnostic_path = RESULTS / "stop_loss_diagnostic_fixed_calendar.csv"
    diagnostic = pd.read_csv(diagnostic_path) if diagnostic_path.exists() else pd.DataFrame()
    return summary, orders, trades, wf, car, hp, scores, diagnostic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), BORDER)


def set_table_width(table, widths):
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)
            set_cell_margins(row.cells[idx])
            row.cells[idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def format_run(run, size=10.5, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def style_paragraph(p, size=10.5, color=None, bold=False, align=None, after=6):
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.1
    if align:
        p.alignment = align
    for run in p.runs:
        format_run(run, size=size, color=color, bold=bold)


def add_h(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    format_run(run, size={1: 16, 2: 13, 3: 12}.get(level, 11), bold=True, color=BLUE if level < 3 else DARK)
    return p


def add_body(doc, text, after=6):
    p = doc.add_paragraph(text)
    style_paragraph(p, after=after)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style=None)
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.15)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run("• ")
        format_run(run, size=10.5, color=DARK, bold=True)
        run = p.add_run(item)
        format_run(run, size=10.5)


def add_callout(doc, label, text):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Inches(6.4)
    set_table_borders(table)
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALLOUT)
    set_cell_margins(cell, top=120, bottom=120, start=180, end=180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(label + "：")
    format_run(r, bold=True, color=DARK)
    r = p.add_run(text)
    format_run(r)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(doc, headers, rows, widths=None, aligns=None, font_size=9.3):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_borders(table)
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_shading(hdr[i], LIGHT)
        set_cell_margins(hdr[i])
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(str(h))
        format_run(run, size=font_size, bold=True, color=DARK)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_margins(cells[i])
            p = cells[i].paragraphs[0]
            if aligns and aligns[i] == "right":
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            elif aligns and aligns[i] == "center":
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(str(value))
            format_run(run, size=font_size)
    if widths:
        set_table_width(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table


def pct(x):
    return f"{float(x) * 100:.2f}%"


def money(x):
    return f"{float(x):,.2f}"


def build():
    summary, orders, trades, wf, car, hp, scores, diagnostic = load_data()

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)

    header = section.header.paragraphs[0]
    header.text = "体育大事件事件驱动量化投资系统 · 中期报告"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    style_paragraph(header, size=9, color=MUTED, after=0)

    footer = section.footer.paragraphs[0]
    footer.text = "QuantInvest · 2026-05-26 数据冻结"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(footer, size=9, color=MUTED, after=0)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    r = title.add_run("体育大事件事件驱动量化投资系统中期报告")
    format_run(r, size=22, bold=True, color=DARK)
    subtitle = doc.add_paragraph("以 2026 FIFA 世界杯为核心模拟交易案例")
    style_paragraph(subtitle, size=12, color=MUTED, after=10)

    add_table(
        doc,
        ["项目", "内容"],
        [
            ["数据冻结日", "2026-05-26 收盘后"],
            ["模拟订单日", "2026-05-27"],
            ["系统版本", summary["code_version"]],
            ["配置 Hash", summary["config_hash"][:16] + "..."],
            ["运行 Hash", summary["run_hash"][:16] + "..."],
        ],
        widths=[1.45, 4.95],
        aligns=["center", "left"],
    )

    add_callout(
        doc,
        "核心结论",
        f"系统以体育大事件库为起点，将事件证据、市场注意力、价格行为、流动性和交易成本统一到自动化流程中。2026-05-26 的真实数据运行结果显示，模型生成的模拟盘目标仓位为 {pct(summary['paper_total_target_weight'])}，实际买入金额为 {money(summary['paper_total_estimated_trade_value'])} 元，预计交易成本为 {money(summary['paper_total_expected_cost'])} 元。"
    )

    add_h(doc, "一、中期要求对照", 1)
    add_table(
        doc,
        ["要求", "系统产物", "状态"],
        [
            ["策略基本思想", "体育大事件事件研究 + 自动组合生成", "完成"],
            ["核心流程", "事件 Agent、证据闭环、因子计算、验证、订单生成", "完成"],
            ["至少一轮回测", "2022 世界杯主回测，另含欧洲杯和奥运会", "完成"],
            ["回测及结果分析", "净收益、成本、止损、CAR、扩展验证", "完成"],
            ["模拟交易配置", "2026-05-26 冻结数据生成 2026-05-27 订单", "完成"],
        ],
        widths=[1.65, 3.7, 1.05],
        aligns=["center", "left", "center"],
    )

    add_h(doc, "二、策略思想", 1)
    add_body(doc, "项目研究的问题不是单纯判断世界杯概念股是否上涨，而是检验确定性体育大事件在资本市场中是否会形成可交易的赛前注意力扩散、产业链重估和事件兑现效应。世界杯是核心案例，系统同时支持欧洲杯、奥运会、亚运会、赛事抽签、赛程公布、门票销售、赞助官宣等体育事件。")
    add_bullets(
        doc,
        [
            "事件确定性：赛事时间和影响范围可由官方渠道确认。",
            "注意力扩散：赛事临近时，搜索、新闻、社媒和成交额常出现同步升温。",
            "产业链暴露：显示设备、IP 衍生品、体育设施、体育运营、传媒彩票等链条受影响程度不同。",
            "兑现效应：赛前预期可能提前反映，临近开幕或开幕后可能出现利好兑现。",
        ],
    )

    add_h(doc, "三、理论基础", 1)
    add_table(
        doc,
        ["模块", "理论来源", "系统实现"],
        [
            ["事件研究", "MacKinlay 事件研究法", "市场模型计算 AR/CAR，剥离大盘影响"],
            ["注意力", "投资者注意力与媒体情绪研究", "成交额冲击作为中期代理，后续接入搜索和社媒"],
            ["动量", "Jegadeesh-Titman 与 Carhart", "20 日价格动量进入横截面排序"],
            ["风格控制", "Fama-French 风格风险", "避免收益只来自小盘或高估值概念暴露"],
            ["流动性", "Amihud 流动性思想", "20 日成交额、滑点分层和整数手约束"],
        ],
        widths=[1.2, 2.25, 2.95],
        aligns=["center", "left", "left"],
    )

    add_h(doc, "四、系统流程", 1)
    add_body(doc, "系统采用自动运行、可审计、可人工配置的流程。研究者可以修改事件池、超参数空间和风控约束，但每次运行都会冻结配置、行情快照、证据输出和订单文件。")
    add_table(
        doc,
        ["步骤", "输入", "输出"],
        [
            ["事件库", "赛事日程、事件类型、官方来源", "event_id、event_date、known_at"],
            ["Agent 抽取", "事件和公司证据", "事件-股票关系 JSON"],
            ["证据闭环", "来源、时间、实体、置信度", "通过校验的暴露库"],
            ["因子计算", "行情快照、暴露、成交额、收益", "横截面因子面板"],
            ["模型评分", "约束 RankIC 权重与超参数", "AlphaScore"],
            ["组合生成", "仓位、行业、成本和 A 股约束", "目标权重与订单"],
            ["回测/模拟盘", "历史事件或当日快照", "绩效、CAR、交易日志"],
        ],
        widths=[1.15, 2.55, 2.7],
        aligns=["center", "left", "left"],
    )

    add_h(doc, "五、Agent 可信闭环", 1)
    add_body(doc, "大模型或规则 Agent 只负责结构化事件和证据，不直接决定买卖。任何事件-公司关系都必须通过 Schema、Source、Time、Entity、Quant 五道门，失败则不能进入因子库。")
    add_table(
        doc,
        ["校验门", "检验内容", "失败处理"],
        [
            ["Schema", "JSON 字段完整、类型合法", "拒绝或重试"],
            ["Source", "每个关系必须有来源和证据文本", "剔除关系"],
            ["Time", "发布时间不得晚于决策日", "防止超前信息"],
            ["Entity", "股票代码、公司名和股票池匹配", "拒绝或人工复核"],
            ["Quant", "置信度为正且通过交易约束", "才可进入评分"],
        ],
        widths=[1.15, 3.0, 2.25],
        aligns=["center", "left", "left"],
    )
    val = summary["validation"]
    add_table(
        doc,
        ["指标", "数值"],
        [
            ["候选关系数", int(val["total"])],
            ["通过数", int(val["accepted"])],
            ["拒绝数", int(val["rejected"])],
            ["通过率", pct(val["accept_rate"])],
            ["幻觉代理率", pct(val["hallucination_proxy_rate"])],
        ],
        widths=[2.2, 1.5],
        aligns=["center", "right"],
    )

    add_h(doc, "六、止损触发归因", 1)
    add_body(doc, "体育事件策略的最大风险不是单笔亏损本身，而是错误地把“事件会发生”理解为“现在就值得交易”。系统使用风险诊断实验识别止损集中出现的场景：一类是建仓窗口过早，事件注意力尚未形成；另一类是事件类型与当前股票池的产业链映射较弱。")
    if not diagnostic.empty:
        diag_rows = []
        for _, row in diagnostic.groupby(["event_id", "exit_reason"]).agg(n=("ticker", "size"), pnl=("net_pnl", "sum")).reset_index().iterrows():
            diag_rows.append([row["event_id"], row["exit_reason"], int(row["n"]), money(row["pnl"])])
        add_table(
            doc,
            ["诊断事件", "退出类型", "交易数", "净 PnL"],
            diag_rows,
            widths=[2.45, 1.25, 0.85, 1.45],
            aligns=["left", "center", "center", "right"],
        )
    add_bullets(
        doc,
        [
            "2022 世界杯的早期窗口中，部分股票在 T-90 左右触发止损，说明事件预热尚未真正进入交易阶段，固定日历建仓容易过早暴露在市场噪声中。",
            "2024 巴黎奥运会在当前 A 股体育候选池中映射较弱，诊断实验中三笔全部止损；事件研究也显示其 T-60 到 T-10 的 CAAR 为负。",
            "固定 8% 止损没有考虑个股波动差异，高波动主题股容易被正常波动洗出；止损规则必须与趋势确认、事件适配度和波动率管理同时使用。",
        ],
    )
    add_body(doc, "Kaminski 和 Lo 的止损研究表明，止损规则是否有效取决于收益过程是否存在趋势和状态转换；Moskowitz、Ooi 和 Pedersen 的时间序列动量研究支持在入场时加入趋势确认；Moreira 和 Muir 的波动率管理思想则说明风险暴露应随波动率变化动态调整。系统因此将固定日历建仓扩展为事件适配度、入场窗口、波动率自适应止损和止损率惩罚共同作用的风险框架。")

    add_h(doc, "七、模型设计", 1)
    add_body(doc, "模型采用符号约束 RankIC。RankIC 用历史事件样本估计因子与未来净超额收益的排序关系；符号约束和理论先验用于防止小样本噪声把事件暴露、动量、流动性等应为正向的因子学成反向权重。事件暴露同时作为交易资格门槛，弱相关股票即使短期成交活跃，也不会自动获得高仓位。")
    weights = summary["factor_weights"]
    add_table(
        doc,
        ["因子", "权重", "含义"],
        [
            ["Exposure", f"{weights['exposure']:.3f}", "事件产业链关系强度"],
            ["Attention", f"{weights['attention']:.3f}", "成交额注意力冲击"],
            ["Momentum", f"{weights['momentum']:.3f}", "20 日价格趋势"],
            ["Liquidity", f"{weights['liquidity']:.3f}", "交易可执行性"],
            ["Low Volatility", f"{weights['low_volatility']:.3f}", "风险控制倾向"],
        ],
        widths=[1.45, 1.0, 3.95],
        aligns=["center", "right", "left"],
    )

    add_body(doc, "风险控制部分由四个模块组成：第一，事件适配度低于门槛时不交易；第二，建仓窗口由超参数搜索决定，而不是默认越早越好；第三，个股止损采用 `max(固定止损, 波动率倍数 * 20 日波动率)`，并设置最大止损上限；第四，超参数选择函数惩罚止损率，使高收益但频繁止损的参数组合不会被优先选择。")

    add_h(doc, "八、超参数分析", 1)
    bp = summary["best_params"]
    add_table(
        doc,
        ["参数", "选择值", "解释"],
        [
            ["attention_lookback", bp["attention_lookback"], "注意力冲击观察窗口"],
            ["momentum_lookback", bp["momentum_lookback"], "价格动量窗口"],
            ["entry_days_before_event", bp["entry_days_before_event"], "赛前建仓窗口"],
            ["exit_days_before_event", bp["exit_days_before_event"], "赛前退出窗口"],
            ["top_k", bp["top_k"], "组合候选数量"],
            ["stop_loss", pct(bp["stop_loss"]), "个股止损线"],
            ["stop_vol_multiplier", bp["stop_vol_multiplier"], "波动率自适应止损倍数"],
            ["max_effective_stop_loss", pct(bp["max_effective_stop_loss"]), "动态止损上限"],
            ["min_exposure_for_trade", bp["min_exposure_for_trade"], "交易暴露门槛"],
            ["min_event_fit_for_trade", bp["min_event_fit_for_trade"], "事件适配度门槛"],
            ["attention_z_cap", bp["attention_z_cap"], "注意力拥挤上限"],
        ],
        widths=[2.1, 1.1, 3.2],
        aligns=["center", "center", "left"],
    )
    add_body(doc, f"本次超参数网格共 {len(hp)} 组，按时间顺序进行扩展窗口验证。系统先用 2022 世界杯训练，再验证 2024 欧洲杯；随后用前两个事件训练，再验证 2024 巴黎奥运会。选择标准同时考虑扩展窗口验证、全样本稳健性、交易成本和止损率，避免只追求单个事件的最高收益。")

    add_h(doc, "九、回测结果", 1)
    b = summary["backtest_summary"]
    w = summary["walk_forward_summary"]
    stop_rate = float((trades["exit_reason"] == "stop_loss").mean()) if not trades.empty else 0.0
    wf_stop_rate = float((wf["exit_reason"] == "stop_loss").mean()) if not wf.empty else 0.0
    add_table(
        doc,
        ["指标", "历史事件回测", "扩展窗口验证"],
        [
            ["交易笔数", int(b["num_trades"]), int(w["num_trades"])],
            ["净收益率", pct(b["net_return"]), pct(w["net_return"])],
            ["平均单笔毛收益", pct(b["gross_return_mean"]), pct(w["gross_return_mean"])],
            ["胜率", pct(b["win_rate"]), pct(w["win_rate"])],
            ["止损率", pct(stop_rate), pct(wf_stop_rate)],
            ["总交易成本", money(b["total_cost"]), money(w["total_cost"])],
            ["成本 / 初始资金", pct(b["cost_to_initial_cash"]), pct(w["cost_to_initial_cash"])],
        ],
        widths=[1.9, 2.15, 2.15],
        aligns=["center", "right", "right"],
    )
    add_body(doc, "历史事件回测和扩展窗口验证均为正，且当前交易集合没有触发止损。巴黎奥运会未进入交易集合，不代表系统不支持奥运会，而是表示当前股票池与该事件的产业链适配度不足；若后续建立专门的奥运会旅游、场馆、体育用品和转播候选池，系统可按同一流程重新评估。")

    event_rows = []
    for _, row in trades.groupby(["event_id", "exit_reason"]).agg(n=("ticker", "size"), pnl=("net_pnl", "sum")).reset_index().iterrows():
        event_rows.append([row["event_id"], row["exit_reason"], int(row["n"]), money(row["pnl"])])
    add_table(
        doc,
        ["事件", "退出类型", "交易数", "净 PnL"],
        event_rows,
        widths=[2.45, 1.25, 0.85, 1.45],
        aligns=["left", "center", "center", "right"],
    )

    add_h(doc, "十、事件研究结果", 1)
    car_rows = []
    for _, row in car.iterrows():
        car_rows.append([row["event_id"], row["window"], pct(row["caar"]), pct(row["positive_rate"])])
    add_table(
        doc,
        ["事件", "窗口", "CAAR", "正 CAR 比例"],
        car_rows,
        widths=[2.35, 1.3, 1.1, 1.35],
        aligns=["left", "center", "right", "right"],
    )
    add_body(doc, "事件研究显示，2022 世界杯在 T-60 到 T-10 和 T-30 到 T-10 窗口具有明显正异常收益，而 T-10 到 T+10 转为负值，符合赛前预热和临近兑现的交易假设。欧洲杯信号较弱，巴黎奥运会与当前 A 股候选池的赛前映射并不稳定。")

    add_h(doc, "十一、2026-05-26 模拟交易配置", 1)
    add_table(
        doc,
        ["项目", "配置"],
        [
            ["初始资金", "1,000,000 元"],
            ["事件", "2026 FIFA 世界杯开幕"],
            ["决策日", "2026-05-26 收盘后"],
            ["订单日", "2026-05-27"],
            ["事件策略仓位上限", "30% 保守兜底"],
            ["系统目标仓位", pct(summary["paper_total_target_weight"])],
            ["预计买入金额", money(summary["paper_total_estimated_trade_value"])],
            ["预计交易成本", money(summary["paper_total_expected_cost"])],
        ],
        widths=[2.0, 4.1],
        aligns=["center", "left"],
    )
    order_rows = []
    for _, row in orders.iterrows():
        if int(row["target_shares"]) <= 0:
            continue
        order_rows.append([
            row["ticker"],
            row["company"],
            pct(row["target_weight"]),
            f"{row['reference_close']:.2f}",
            int(row["target_shares"]),
            money(row["estimated_trade_value"]),
            money(row["expected_cost"]),
        ])
    add_table(
        doc,
        ["股票", "公司", "权重", "收盘价", "股数", "金额", "成本"],
        order_rows,
        widths=[1.05, 1.0, 0.8, 0.8, 0.85, 1.05, 0.85],
        aligns=["center", "center", "right", "right", "right", "right", "right"],
        font_size=8.8,
    )
    add_body(doc, "模拟盘没有满仓追入，实际订单集中在海信视像、共创草坪和中体产业。原因是系统在临近开幕窗口引入了注意力拥挤上限、事件暴露门槛和低波动约束，弱相关或风险贡献较差的股票会被降权或拒单。")

    add_h(doc, "十二、市场约束与成本", 1)
    add_table(
        doc,
        ["约束", "系统处理"],
        [
            ["A 股 T+1", "回测持仓跨日，模拟订单按下一交易日执行"],
            ["100 股整数手", "目标股数向下取整，不足一手自动拒单"],
            ["涨跌停", "主板 10%、创业板 20% 规则纳入约束"],
            ["佣金", "双边 0.03%，最低 5 元"],
            ["印花税", "卖出单边 0.05%"],
            ["滑点", "按 20 日成交额分层估计"],
            ["仓位", "单股 12%、行业 25%、事件兜底 30%"],
        ],
        widths=[1.5, 4.9],
        aligns=["center", "left"],
    )

    add_h(doc, "十三、风险与后续计划", 1)
    add_bullets(
        doc,
        [
            "样本量仍有限，结项前需要加入更多体育事件和事件前置节点。",
            "当前注意力因子使用成交额冲击代理，后续接入搜索指数、新闻数量和社媒讨论量。",
            "Agent 证据集需要扩展为人工标注样本，评估准确率、召回率和幻觉率。",
            "模拟盘从 2026-05-27 开始记录每日持仓、净值、交易成本、止盈止损和信号变化。",
            "结项报告需要回顾模拟持仓效果，解释收益或亏损来自事件、因子、成本还是风控。",
        ],
    )

    add_h(doc, "十四、产物索引", 1)
    add_table(
        doc,
        ["产物", "路径"],
        [
            ["系统代码", "quant_sports_event/"],
            ["配置", "configs/strategy.json, configs/events.json, configs/universe.json"],
            ["测试", "tests/test_system.py"],
            ["行情快照", "outputs/data_snapshots/2026-05-26/"],
            ["回测交易", "outputs/results/backtest_trades_2026-05-26.csv"],
            ["扩展验证", "outputs/results/walk_forward_trades_2026-05-26.csv"],
            ["模拟订单", "outputs/results/paper_orders_2026-05-27.csv"],
            ["运行摘要", "outputs/results/run_summary_2026-05-26.json"],
        ],
        widths=[1.4, 5.0],
        aligns=["center", "left"],
        font_size=8.6,
    )

    add_h(doc, "十五、参考文献与来源", 1)
    refs = [
        "MacKinlay, A. C. (1997). Event Studies in Economics and Finance.",
        "Edmans, A., García, D., & Norli, Ø. (2007). Sports Sentiment and Stock Returns.",
        "Tetlock, P. C. (2007). Giving Content to Investor Sentiment.",
        "Jegadeesh, N., & Titman, S. (1993). Returns to Buying Winners and Selling Losers.",
        "Fama, E. F., & French, K. R. (1993). Common Risk Factors in the Returns on Stocks and Bonds.",
        "Carhart, M. M. (1997). On Persistence in Mutual Fund Performance.",
        "Kaminski, K. M., & Lo, A. W. (2014). When Do Stop-Loss Rules Stop Losses?",
        "Moskowitz, T. J., Ooi, Y. H., & Pedersen, L. H. (2012). Time Series Momentum.",
        "Moreira, A., & Muir, T. (2017). Volatility-Managed Portfolios.",
        "Yao et al. (2023). ReAct: Synergizing Reasoning and Acting in Language Models.",
        "Gou et al. (2024). CRITIC: Large Language Models Can Self-Correct with Tool-Interactive Critiquing.",
        "Shanghai Stock Exchange. Stock Trading Mechanism.",
        "FIFA World Cup 2026 Match Schedule.",
    ]
    add_bullets(doc, refs)

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build())
